import os
import re
import shutil
import pythoncom
import win32com.client
from collections import defaultdict
from database import get_base_dir


def _cell_replace_checkbox(cell, option, checked=True):
    """Replace a checkbox in a Word table cell while preserving paragraph structure.

    Each paragraph's Range.Text ends with chr(13) (para mark). We strip only chr(7)
    (cell terminator on the last para), do the substitution, then write the text
    back including the chr(13) so the paragraph mark is never destroyed.
    """
    mark = '☑' if checked else '□'
    cell_rng = cell.Range
    for i in range(1, cell_rng.Paragraphs.Count + 1):
        p = cell_rng.Paragraphs(i)
        raw = p.Range.Text          # includes trailing \r (and \x07 on last para)
        clean = raw.replace('\x07', '')  # remove cell terminator, keep \r
        if ('□' + option) not in clean:
            continue
        new_txt = clean.replace('□' + option, mark + option)
        p.Range.Text = new_txt      # \r is still there, para mark survives
        break

def _set_issues_cell(cell, issues):
    """Write numbered issues separated by semicolons on a single line."""
    rng = cell.Range
    rng.MoveEnd(1, -1)
    if issues:
        parts = [f'{idx+1}.{str(v)}' for idx, v in enumerate(issues)]
        rng.Text = '；'.join(parts)
    else:
        rng.Text = '无'






def _insert_sign_img(doc, names, sign_data):
    if not sign_data: return
    
    import uuid, base64, os
    abs_path = None
    if sign_data.startswith("data:image"):
        head, b64_str = sign_data.split(',', 1)
        img_data = base64.b64decode(b64_str)
        tmp_path = os.path.abspath(f"uploads/signatures/tmp_{uuid.uuid4().hex}.png")
        os.makedirs("uploads/signatures", exist_ok=True)
        with open(tmp_path, "wb") as f:
            f.write(img_data)
        abs_path = tmp_path
    elif "file=" in sign_data:
        rel_path = sign_data.split("file=")[1]
        abs_path = os.path.abspath(rel_path)
        
    if not abs_path or not os.path.exists(abs_path): return
    
    for name in names:
        if doc.Bookmarks.Exists(name):
            rng = doc.Bookmarks(name).Range
            rng.Text = ""
            shape = rng.InlineShapes.AddPicture(FileName=abs_path, LinkToFile=False, SaveWithDocument=True)
            ratio = shape.Width / shape.Height if shape.Height else 2
            shape.Height = 30
            shape.Width = 30 * ratio

from datetime import datetime

def replace_common_bookmarks(doc, county_name: str, date_str: str = None):
    """
    通用替换所有附件中的全局书签：
    1. county_name -> 县级区域名称
    2. date -> 当前时间 XXXX年XX月XX日
    """
    if not date_str:
        now = datetime.now()
        date_str = f"{now.year}年{now.month:02d}月{now.day:02d}日"

    if doc.Bookmarks.Exists("county_name"):
        try:
            doc.Bookmarks("county_name").Range.Text = county_name
        except Exception as e:
            print(f"替换书签 county_name 异常: {e}")

    if doc.Bookmarks.Exists("county_name1"):
        try:
            doc.Bookmarks("county_name1").Range.Text = county_name
        except Exception as e:
            print(f"替换书签 county_name1 异常: {e}")

    if doc.Bookmarks.Exists("date"):
        try:
            doc.Bookmarks("date").Range.Text = date_str
        except Exception as e:
            print(f"替换书签 date 异常: {e}")

def replace_att6_bookmarks(doc, county_name: str, date_str: str = None):
    """
    针对附件6各分表书签替换：
    county_name1/2/3/4 -> 县级区域名称
    date1/2/3/4 -> 当前时间 XXXX年XX月XX日
    """
    if not date_str:
        now = datetime.now()
        date_str = f"{now.year}年{now.month:02d}月{now.day:02d}日"

    for i in range(1, 5):
        bm_c = f"county_name{i}"
        if doc.Bookmarks.Exists(bm_c):
            try:
                doc.Bookmarks(bm_c).Range.Text = county_name
            except Exception as e:
                print(f"替换书签 {bm_c} 异常: {e}")

        bm_d = f"date{i}"
        if doc.Bookmarks.Exists(bm_d):
            try:
                doc.Bookmarks(bm_d).Range.Text = date_str
            except Exception as e:
                print(f"替换书签 {bm_d} 异常: {e}")

def _fill_bookmarks(doc, names, text):
    """Fill named bookmarks with text, padding with spaces to original length."""
    for name in names:
        if doc.Bookmarks.Exists(name):
            bm = doc.Bookmarks(name)
            bm_len = bm.End - bm.Start
            padded = (text + " " * bm_len)[:bm_len]
            bm_rng = bm.Range
            bm_rng.Text = padded
            # Re-add bookmark since .Text assignment removes it
            doc.Bookmarks.Add(name, bm_rng)

def sanitize_filename(name):
    clean = re.sub(r'[\\/*?:"<>|（）() ]', "", str(name or "")).strip()
    return clean if clean else "组"

def format_score(score):
    if score is None:
        return ""
    if isinstance(score, (int, float)):
        if score == int(score):
            return str(int(score))
        return f"{score:.1f}"
    return str(score)

def calculate_neiye_subscores(form_data):
    # 1. 机制运行 (满分15)
    d_m1 = 2.0 if form_data.get('mech_1') else 0.0
    d_m2 = 0.0
    for opt in form_data.get('mech_2', []):
        if '支付不规范' in opt: d_m2 += 4.0
        elif '支付不及时' in opt: d_m2 += 4.0
        elif '兜底' in opt: d_m2 += 2.0
    d_m2 = min(d_m2, 10.0)
    d_m3 = 2.0 if form_data.get('mech_3') else 0.0
    
    # 培训 (可为列表或count)
    if isinstance(form_data.get('mech_4'), list):
        d_m4 = min(len(form_data.get('mech_4', [])) * 0.5, 1.0)
    else:
        d_m4 = min(float(form_data.get('mech_4_count', 0) or 0) * 0.5, 1.0)
    
    deduct_mech = min(d_m1 + d_m2 + d_m3 + d_m4, 15.0)
    score_mech = max(15.0 - deduct_mech, 0.0)
    
    # 2. 程序规范 (满分30)
    d_p1 = 5.0 if form_data.get('prog_1') else 0.0
    d_p2 = min(len(form_data.get('prog_2', [])) * 0.5, 5.0)
    d_p3 = 5.0 if form_data.get('prog_3') else 0.0
    
    p4_list = form_data.get('prog_4', [])
    if any('没有公示材料' in x or '不足15天' in x for x in p4_list):
        d_p4 = 2.0
    else:
        d_p4 = min(len(p4_list) * 0.5, 2.0)
        
    d_p5 = 3.0 if form_data.get('prog_5') else 0.0
    d_p6 = 5.0 if form_data.get('prog_6') else 0.0
    d_p7 = 5.0 if form_data.get('prog_7') else 0.0
    
    deduct_prog = min(d_p1 + d_p2 + d_p3 + d_p4 + d_p5 + d_p6 + d_p7, 30.0)
    score_prog = max(30.0 - deduct_prog, 0.0)
    
    # 3. 政策落实 (满分15)
    d_pol1 = min(len(form_data.get('policy_1', [])) * 1.0, 3.0)
    c_2_1 = float(form_data.get('policy_2_1', 0) or 0)
    c_2_2 = float(form_data.get('policy_2_2', 0) or 0)
    d_pol2 = min((c_2_1 + c_2_2) * 1.0, 3.0)
    
    c_3_1 = float(form_data.get('policy_3_1', 0) or 0)
    c_3_2 = float(form_data.get('policy_3_2', 0) or 0)
    d_pol3 = min((c_3_1 + c_3_2) * 1.0, 3.0)
    
    d_pol4 = min(len(form_data.get('policy_4', [])) * 0.5, 3.0)
    d_pol5 = min(len(form_data.get('policy_5', [])) * 0.5, 3.0)
    
    deduct_policy = min(d_pol1 + d_pol2 + d_pol3 + d_pol4 + d_pol5, 15.0)
    score_policy = max(15.0 - deduct_policy, 0.0)
    
    # 4. 工作成效 (满分10)
    deduct_effect = min(len(form_data.get('effect_1', [])) * 1.0, 10.0)
    score_effect = max(10.0 - deduct_effect, 0.0)
    
    total_score = score_mech + score_prog + score_policy + score_effect
    
    return {
        "deduct": {
            "m1": d_m1, "m2": d_m2, "m3": d_m3, "m4": d_m4, "mech": deduct_mech,
            "p1": d_p1, "p2": d_p2, "p3": d_p3, "p4": d_p4, "p5": d_p5, "p6": d_p6, "p7": d_p7, "prog": deduct_prog,
            "pol1": d_pol1, "pol2": d_pol2, "pol3": d_pol3, "pol4": d_pol4, "pol5": d_pol5, "policy": deduct_policy,
            "effect": deduct_effect
        },
        "score": {
            "mech": score_mech,
            "prog": score_prog,
            "policy": score_policy,
            "effect": score_effect,
            "total": total_score
        }
    }

def fill_table_1(t1, form_data, scores):
    # R2: 配套延包
    if any("未制定方案" in x for x in form_data.get("mech_1", [])):
        _cell_replace_checkbox(t1.Rows(2).Cells(4), "未制定方案")
    if any("直接套用上级方案" in x for x in form_data.get("mech_1", [])):
        _cell_replace_checkbox(t1.Rows(2).Cells(4), "直接套用上级方案")
    if any("制定程序不合法" in x for x in form_data.get("mech_1", [])):
        _cell_replace_checkbox(t1.Rows(2).Cells(4), "制定程序不合法")
    t1.Rows(2).Cells(6).Range.Text = f"{scores['deduct']['m1']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['m1'] > 0 else "0"
    
    # R3: 经费保障
    if any("支付不规范" in x for x in form_data.get("mech_2", [])):
        _cell_replace_checkbox(t1.Rows(3).Cells(4), "支付不规范")
    if any("支付不及时" in x for x in form_data.get("mech_2", [])):
        _cell_replace_checkbox(t1.Rows(3).Cells(4), "支付不及时")
    if any("经费没有县级兜底" in x for x in form_data.get("mech_2", [])):
        _cell_replace_checkbox(t1.Rows(3).Cells(4), "经费没有县级兜底")
    t1.Rows(3).Cells(6).Range.Text = f"{scores['deduct']['m2']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['m2'] > 0 else "0"

    # R4: 宣传
    if any("没有宣传材料" in x for x in form_data.get("mech_3", [])):
        _cell_replace_checkbox(t1.Rows(4).Cells(4), "没有宣传材料")
    t1.Rows(4).Cells(6).Range.Text = f"{scores['deduct']['m3']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['m3'] > 0 else "0"

    # R5: 培训
    if any("没有培训材料" in x for x in form_data.get("mech_4", [])):
        _cell_replace_checkbox(t1.Rows(5).Cells(4), "没有培训材料")
    if any("没有分批次培训" in x for x in form_data.get("mech_4", [])):
        _cell_replace_checkbox(t1.Rows(5).Cells(4), "没有分批次培训")
    if any("培训材料不齐全" in x for x in form_data.get("mech_4", [])):
        _cell_replace_checkbox(t1.Rows(5).Cells(4), "培训材料不齐全")
    if any("培训未覆盖县乡村组" in x for x in form_data.get("mech_4", [])):
        _cell_replace_checkbox(t1.Rows(5).Cells(4), "培训未覆盖县乡村组")
    t1.Rows(5).Cells(6).Range.Text = f"{scores['deduct']['m4']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['m4'] > 0 else "0"

    # R6: 总计扣分
    t1.Rows(6).Cells(3).Range.Text = f"{scores['deduct']['mech']:.1f}".rstrip('0').rstrip('.')
    
    # R7: 重要问题描述
    issues = []
    if form_data.get("mech_1"): issues.extend(form_data.get("mech_1", []))
    if form_data.get("mech_2"): issues.extend(form_data.get("mech_2", []))
    if form_data.get("mech_3"): issues.extend(form_data.get("mech_3", []))
    if form_data.get("mech_4"): issues.extend(form_data.get("mech_4", []))
    _set_issues_cell(t1.Rows(7).Cells(2), issues) 

def fill_table_2(t2, form_data, scores):
    # R2: 成立机构
    if any("未召开会议" in x for x in form_data.get("prog_1", [])):
        _cell_replace_checkbox(t2.Rows(2).Cells(4), "未召开会议")
    if any("未公示工作组名单" in x for x in form_data.get("prog_1", [])):
        _cell_replace_checkbox(t2.Rows(2).Cells(4), "未公示工作组名单")
    if any("参会人数不足法定数量" in x for x in form_data.get("prog_1", [])):
        _cell_replace_checkbox(t2.Rows(2).Cells(4), "参会人数不足法定数量")
    t2.Rows(2).Cells(6).Range.Text = f"{scores['deduct']['p1']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p1'] > 0 else "0"

    # R3: 摸底核实
    if any("没有进行摸底" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "没有进行摸底")
    if any("摸底表农户未签署" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "摸底表农户未签署")
    if any("摸底表中没有表达延包意愿" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "摸底表中没有表达延包意愿")
    if any("摸底表其它签署不齐全" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "摸底表其它签署不齐全")
    if any("特殊人员摸底不清或未统计" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "特殊人员摸底不清或未统计")
    if any("户变化未统计" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "户变化未统计")
    if any("矛盾纠纷未登记或处理不当" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "矛盾纠纷未登记或处理不当")
    if any("承包地变化未摸清" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "承包地变化未摸清")
    if any("没有应确尽确" in x for x in form_data.get("prog_2", [])):
        _cell_replace_checkbox(t2.Rows(3).Cells(4), "没有应确尽确")
    t2.Rows(3).Cells(6).Range.Text = f"{scores['deduct']['p2']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p2'] > 0 else "0"

    # R4: 制定方案
    if any("没有延包方案" in x for x in form_data.get("prog_3", [])):
        _cell_replace_checkbox(t2.Rows(4).Cells(4), "没有延包方案")
    if any("延包方案未上报" in x for x in form_data.get("prog_3", [])):
        _cell_replace_checkbox(t2.Rows(4).Cells(4), "延包方案未上报")
    if any("延包方案未公示" in x for x in form_data.get("prog_3", [])):
        _cell_replace_checkbox(t2.Rows(4).Cells(4), "延包方案未公示")
    if any("未召开会议讨论延包方案" in x for x in form_data.get("prog_3", [])):
        _cell_replace_checkbox(t2.Rows(4).Cells(4), "未召开会议讨论延包方案")
    t2.Rows(4).Cells(6).Range.Text = f"{scores['deduct']['p3']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p3'] > 0 else "0"

    # R5: 调查公示
    if any("没有公示材料" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "没有公示材料")
    if any("公示不足15天" in x or "没有公示或不足15天" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "公示不足15天")
    if any("公示结果未确认" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "公示结果未确认")
    if any("各类资料不齐全" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "各类资料不齐全")
    if any("各类资料制作粗糙" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "各类资料制作粗糙")
    if any("各类资料签署不规范" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "各类资料签署不规范")
    if any("权属证明材料不齐全" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "权属证明材料不齐全")
    if any("其它证明材料不齐全" in x for x in form_data.get("prog_4", [])):
        _cell_replace_checkbox(t2.Rows(5).Cells(4), "其它证明材料不齐全")
    t2.Rows(5).Cells(6).Range.Text = f"{scores['deduct']['p4']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p4'] > 0 else "0"

    # R6: 签订合同
    if any("合同版本格式不正确" in x for x in form_data.get("prog_5", [])):
        _cell_replace_checkbox(t2.Rows(6).Cells(4), "合同版本格式不正确")
    if any("合同网签率未达到95%" in x for x in form_data.get("prog_5", [])):
        _cell_replace_checkbox(t2.Rows(6).Cells(4), "合同网签率未达到95%")
    if any("没有地块示意图" in x for x in form_data.get("prog_5", [])):
        _cell_replace_checkbox(t2.Rows(6).Cells(4), "没有地块示意图")
    t2.Rows(6).Cells(6).Range.Text = f"{scores['deduct']['p5']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p5'] > 0 else "0"

    # R7: 完善证书
    if any("未进行信息共享" in x for x in form_data.get("prog_6", [])):
        _cell_replace_checkbox(t2.Rows(7).Cells(4), "未进行信息共享")
    if any("未与不动产登记部门有序衔接" in x for x in form_data.get("prog_6", [])):
        _cell_replace_checkbox(t2.Rows(7).Cells(4), "未与不动产登记部门有序衔接")
    t2.Rows(7).Cells(6).Range.Text = f"{scores['deduct']['p6']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p6'] > 0 else "0"

    # R8: 资料归档
    if any("档案整理第三方无涉密档案整理资质" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "档案整理第三方无涉密档案整理资质")
    if any("档案资料收集不齐全" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "档案资料收集不齐全")
    if any("档案分类不符合要求" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "档案分类不符合要求")
    if any("档号与归档章不一致" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "档号与归档章不一致")
    if any("档案变数字化瑕疵" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "档案变数字化瑕疵")
    if any("姓名著录错误" in x for x in form_data.get("prog_7", [])):
        _cell_replace_checkbox(t2.Rows(8).Cells(4), "“一户一档”姓名著录错误")
    t2.Rows(8).Cells(6).Range.Text = f"{scores['deduct']['p7']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['p7'] > 0 else "0"

    # R9: 总计扣分
    t2.Rows(9).Cells(3).Range.Text = f"{scores['deduct']['prog']:.1f}".rstrip('0').rstrip('.')
    
    # R10: 重要问题描述
    issues = []
    if form_data.get("prog_1"): issues.extend(form_data.get("prog_1", []))
    if form_data.get("prog_2"): issues.extend(form_data.get("prog_2", []))
    if form_data.get("prog_3"): issues.extend(form_data.get("prog_3", []))
    if form_data.get("prog_4"): issues.extend(form_data.get("prog_4", []))
    if form_data.get("prog_5"): issues.extend(form_data.get("prog_5", []))
    if form_data.get("prog_6"): issues.extend(form_data.get("prog_6", []))
    if form_data.get("prog_7"): issues.extend(form_data.get("prog_7", []))
    _set_issues_cell(t2.Rows(10).Cells(2), issues) 

def fill_table_3(t3, form_data, scores):
    # R2: 大稳定、小调整
    if any("小调整比率过大或手续不齐全" in x for x in form_data.get("policy_1", [])):
        _cell_replace_checkbox(t3.Rows(2).Cells(4), "小调整比率过大或手续不齐全")
    if any("打乱重分" in x for x in form_data.get("policy_1", [])):
        _cell_replace_checkbox(t3.Rows(2).Cells(4), "打乱重分")
    if any("违法调整或收回承包地" in x for x in form_data.get("policy_1", [])):
        _cell_replace_checkbox(t3.Rows(2).Cells(4), "违法调整或收回承包地")
    t3.Rows(2).Cells(6).Range.Text = f"{scores['deduct']['pol1']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['pol1'] > 0 else "0"

    # R3: 保障土地承包权益 (带数量)
    c_2_1 = int(form_data.get("policy_2_1", 0) or 0)
    c_2_2 = int(form_data.get("policy_2_2", 0) or 0)
    if c_2_1 > 0:
        _cell_replace_checkbox(t3.Rows(3).Cells(4), "未保障特殊群体权益")
        # also append count - use find/replace for the full text
        rng3 = t3.Rows(3).Cells(4).Range
        f3 = rng3.Find; f3.ClearFormatting(); f3.Replacement.ClearFormatting()
        f3.Execute(FindText="未保障特殊群体权益", ReplaceWith=f"未保障特殊群体权益（{c_2_1}起）", Replace=2, MatchCase=True)
    if c_2_2 > 0:
        _cell_replace_checkbox(t3.Rows(3).Cells(4), "未保障无地户权益")
        rng3b = t3.Rows(3).Cells(4).Range
        f3b = rng3b.Find; f3b.ClearFormatting(); f3b.Replacement.ClearFormatting()
        f3b.Execute(FindText="未保障无地户权益", ReplaceWith=f"未保障无地户权益（{c_2_2}起）", Replace=2, MatchCase=True)
    t3.Rows(3).Cells(6).Range.Text = f"{scores['deduct']['pol2']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['pol2'] > 0 else "0"

    # R4: 依法收回消亡户承包地 (带数量)
    c_3_1 = int(form_data.get("policy_3_1", 0) or 0)
    c_3_2 = int(form_data.get("policy_3_2", 0) or 0)
    if c_3_1 > 0:
        _cell_replace_checkbox(t3.Rows(4).Cells(4), "没有应收尽收")
        rng4 = t3.Rows(4).Cells(4).Range
        f4 = rng4.Find; f4.ClearFormatting(); f4.Replacement.ClearFormatting()
        f4.Execute(FindText="没有应收尽收", ReplaceWith=f"没有应收尽收（{c_3_1}起）", Replace=2, MatchCase=True)
    if c_3_2 > 0:
        _cell_replace_checkbox(t3.Rows(4).Cells(4), "采用不正当方式隐匿消亡户")
        rng4b = t3.Rows(4).Cells(4).Range
        f4b = rng4b.Find; f4b.ClearFormatting(); f4b.Replacement.ClearFormatting()
        f4b.Execute(FindText="采用不正当方式隐匿消亡户", ReplaceWith=f"采用不正当方式隐匿消亡户（{c_3_2}起）", Replace=2, MatchCase=True)
    t3.Rows(4).Cells(6).Range.Text = f"{scores['deduct']['pol3']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['pol3'] > 0 else "0"

    # R5: 严格机动地和新增耕地管理
    if any("机动地、新增耕地处置不当" in x for x in form_data.get("policy_4", [])):
        _cell_replace_checkbox(t3.Rows(5).Cells(4), "机动地、新增耕地处置不当")
    if any("机动地比率过高" in x for x in form_data.get("policy_4", [])):
        _cell_replace_checkbox(t3.Rows(5).Cells(4), "机动地比率过高")
    t3.Rows(5).Cells(6).Range.Text = f"{scores['deduct']['pol4']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['pol4'] > 0 else "0"

    # R6: 从严掌握确权确股不确地
    if any("违背农户意愿强行推进" in x for x in form_data.get("policy_5", [])):
        _cell_replace_checkbox(t3.Rows(6).Cells(4), "违背农户意愿强行推进")
    if any("确权确股不确地手续不齐全" in x for x in form_data.get("policy_5", [])):
        _cell_replace_checkbox(t3.Rows(6).Cells(4), "确权确股不确地手续不齐全")
    t3.Rows(6).Cells(6).Range.Text = f"{scores['deduct']['pol5']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['pol5'] > 0 else "0"

    # R7: 总计扣分
    t3.Rows(7).Cells(3).Range.Text = f"{scores['deduct']['policy']:.1f}".rstrip('0').rstrip('.')
    
    # R8: 重要问题描述
    issues = []
    if form_data.get("policy_1"): issues.extend(form_data.get("policy_1", []))
    if c_2_1 > 0: issues.append(f"未保障特殊群体权益（{c_2_1}起）")
    if c_2_2 > 0: issues.append(f"未保障无地户权益（{c_2_2}起）")
    if c_3_1 > 0: issues.append(f"消亡户未应收尽收（{c_3_1}起）")
    if c_3_2 > 0: issues.append(f"不正当方式隐匿消亡户（{c_3_2}起）")
    if form_data.get("policy_4"): issues.extend(form_data.get("policy_4", []))
    if form_data.get("policy_5"): issues.extend(form_data.get("policy_5", []))
    _set_issues_cell(t3.Rows(8).Cells(2), issues) 

def fill_table_4(t4, form_data, scores):
    # R2: 加强风险防范
    if any("未建立矛盾纠纷处置机制" in x for x in form_data.get("effect_1", [])):
        _cell_replace_checkbox(t4.Rows(2).Cells(4), "未建立矛盾纠纷处置机制")
    if any("未建立舆情处置办法" in x for x in form_data.get("effect_1", [])):
        _cell_replace_checkbox(t4.Rows(2).Cells(4), "未建立舆情处置办法")
    if any("没有矛盾纠纷处理台账" in x for x in form_data.get("effect_1", [])):
        _cell_replace_checkbox(t4.Rows(2).Cells(4), "没有矛盾纠纷处理台账")
    t4.Rows(2).Cells(6).Range.Text = f"{scores['deduct']['effect']:.1f}".rstrip('0').rstrip('.') if scores['deduct']['effect'] > 0 else "0"

    # R3: 总计扣分
    t4.Rows(3).Cells(3).Range.Text = f"{scores['deduct']['effect']:.1f}".rstrip('0').rstrip('.')
    
    # R4: 重要问题描述
    issues = []
    if form_data.get("effect_1"): issues.extend(form_data.get("effect_1", []))
    _set_issues_cell(t4.Rows(4).Cells(2), issues) 

def export_neiye_att6_township(qsdwmc, form_data, township_name=None, village_name=None):
    pythoncom.CoInitialize()
    try:
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件6.doc")
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        clean_target = sanitize_filename(qsdwmc)
        out_filename = f"附件6_{county_name}县级自查内业组检查记录表_{clean_target}.doc"
        out_path = os.path.join(base_dir, "backend", "downloads", out_filename)
        shutil.copy(tpl, out_path)
        
        doc = word.Documents.Open(out_path)
        
        # 替换通用与附件6书签：county_name1/2/3/4 -> 县级区域名称，date1/2/3/4 -> 当前时间
        replace_common_bookmarks(doc, county_name)
        replace_att6_bookmarks(doc, county_name)

        # 映射 xzqh_1 ~ xzqh_4 书签：
        # - 若为乡镇级（无村名或村名为镇名），映射为 XX县XX镇
        # - 若为村级，映射为 XX县XX镇XX村
        t_part = township_name or ""
        v_part = village_name or ""
        if not t_part and not v_part:
            # 尝试从 qsdwmc 解析 (例如 "明光街道 - 蔬菜村" 或 "襄河镇")
            if " - " in qsdwmc:
                parts = qsdwmc.split(" - ", 1)
                t_part, v_part = parts[0].strip(), parts[1].strip()
            else:
                t_part = qsdwmc.strip()
                v_part = ""

        # 如果村名与镇名相同或村名为空，判定为乡镇级
        if v_part == t_part:
            v_part = ""

        # 规整县、镇、村文本，防止出现重复前缀（如 明光市明光街道蔬菜村）
        if t_part:
            if t_part.startswith(county_name):
                base_tv = t_part
            else:
                base_tv = f"{county_name}{t_part}"
        else:
            base_tv = county_name

        if v_part:
            if v_part.startswith(base_tv):
                _xzqh_text = v_part
            elif t_part and v_part.startswith(t_part):
                _xzqh_text = f"{county_name}{v_part}" if not v_part.startswith(county_name) else v_part
            else:
                _xzqh_text = f"{base_tv}{v_part}"
        else:
            # 严格满足要求：乡镇级映射为 XX县XX镇
            _xzqh_text = base_tv

        _fill_bookmarks(doc, ["xzqh_1", "xzqh_2", "xzqh_3", "xzqh_4", "xzqh1", "xzqh2", "xzqh3", "xzqh4"], _xzqh_text)
        jcz_sign = form_data.get("jcz_sign")
        fhz_sign = form_data.get("fhz_sign")
        _fill_bookmarks(doc, ["jcz1", "jcz2", "jcz3", "jcz4"], form_data.get("jcz_name") or "")
        _fill_bookmarks(doc, ["fhz1", "fhz2", "fhz3", "fhz4"], form_data.get("fhz_name") or "")
        
        scores = calculate_neiye_subscores(form_data)
        
        fill_table_1(doc.Tables(1), form_data, scores)
        fill_table_2(doc.Tables(2), form_data, scores)
        fill_table_3(doc.Tables(3), form_data, scores)
        fill_table_4(doc.Tables(4), form_data, scores)
        
        doc.Save()
        doc.Close(False)
        word.Quit()
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_neiye_att6_township error:", e)
        try: word.Quit()
        except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_neiye_att6_county(form_data):
    """县级自查：仅导出机制运行 1/4"""
    return export_neiye_att6_mechanism_only(form_data, is_county=True)

def export_neiye_att6_mechanism_only(form_data, is_county=False, township_name=None):
    """
    导出仅包含【机制运行 (1/4)】的附件6检查记录表：
    - is_county=True 时：适用于县级自查，xzqh_1 书签映射为 XX县，输出 附件6_{county_name}县级自查内业组检查记录表（1_4）.doc
    - is_county=False 时：适用于乡镇级自查，xzqh_1 书签映射为 XX县XX镇，输出 附件6_{county_name}县级自查内业组检查记录表_{township_name}（1_4）.doc
    - 自动填充 Table 1 机制运行扣分项与凭证
    - 物理删除 Table 4, 3, 2，仅保留第 1/4 页完整格式并保存
    """
    pythoncom.CoInitialize()
    try:
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件6.doc")
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        
        if is_county or not township_name:
            out_filename = f"附件6_{county_name}县级自查内业组检查记录表（1_4）.doc"
            xzqh_text = county_name
        else:
            clean_ts = sanitize_filename(township_name)
            out_filename = f"附件6_{county_name}县级自查内业组检查记录表_{clean_ts}（1_4）.doc"
            # 严格映射为 XX县XX镇
            if clean_ts.startswith(county_name):
                xzqh_text = clean_ts
            else:
                xzqh_text = f"{county_name}{clean_ts}"

        out_path = os.path.join(base_dir, "backend", "downloads", out_filename)
        shutil.copy(tpl, out_path)
        
        doc = word.Documents.Open(out_path)
        
        # 替换通用与附件6书签：county_name1/2/3/4 -> 县级区域名称，date1/2/3/4 -> 当前时间
        replace_common_bookmarks(doc, county_name)
        replace_att6_bookmarks(doc, county_name)

        # Fill 行政区划名称 bookmark for page 1 (机制运行 1/4)
        _fill_bookmarks(doc, ["xzqh_1", "xzqh1"], xzqh_text)
        jcz_sign = form_data.get("jcz_sign")
        fhz_sign = form_data.get("fhz_sign")
        _fill_bookmarks(doc, ["jcz1"], form_data.get("jcz_name") or "")
        _fill_bookmarks(doc, ["fhz1"], form_data.get("fhz_name") or "")
        
        scores = calculate_neiye_subscores(form_data)
        fill_table_1(doc.Tables(1), form_data, scores)
        
        # Delete Table 4, 3, 2 in reverse order
        while doc.Tables.Count > 1:
            doc.Tables(doc.Tables.Count).Delete()
        
        # Trim paragraphs after Table 1
        end_pos = doc.Content.End
        for p in doc.Paragraphs:
            if "复核者：" in p.Range.Text:
                end_pos = p.Range.End
                break
        if end_pos < doc.Content.End:
            rng = doc.Range(end_pos, doc.Content.End)
            rng.Delete()
        
        doc.Save()
        doc.Close(False)
        word.Quit()
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_neiye_att6_mechanism_only error:", e)
        try: word.Quit()
        except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_neiye_att7(records_by_qsdwdm, township_villages_map=None):
    pythoncom.CoInitialize()
    try:
        from database import get_county_and_townships_sync
        county_info, township_list_raw = get_county_and_townships_sync()
        county_code = county_info.get("code", "341124")
        county_name = county_info.get("name", "全椒县")

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = False
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件7.doc")
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        out_filename = f"附件7_{county_name}县级自查内业组检查得分表.doc"
        out_path = os.path.join(base_dir, "backend", "downloads", out_filename)
        shutil.copy(tpl, out_path)
        
        doc = word.Documents.Open(out_path)
        replace_common_bookmarks(doc, county_name)
        t = doc.Tables(1)
        
        # 1. County row (Row 2)
        county_rec = records_by_qsdwdm.get(county_code) or records_by_qsdwdm.get("341124")
        if county_rec:
            c_scores = calculate_neiye_subscores(county_rec.get("form_data", {}))
            county_mech = c_scores['score']['mech']
        else:
            county_mech = 15.0

        t.Rows(2).Cells(1).Range.Text = "1"
        t.Rows(2).Cells(2).Range.Text = county_name
        t.Rows(2).Cells(3).Range.Text = f"{county_mech:.1f}".rstrip('0').rstrip('.')
        t.Rows(2).Cells(4).Range.Text = "/"
        t.Rows(2).Cells(5).Range.Text = "/"
        t.Rows(2).Cells(6).Range.Text = "/"
        t.Rows(2).Cells(7).Range.Text = "/"

        township_list = [(t["code"], t["name"]) for t in township_list_raw]
        N = len(township_list)
        # 表格初始结构：Row 1 表头，Row 2 县级行，最后一行是总评行
        current_reserved_ts_rows = t.Rows.Count - 3
        if N > current_reserved_ts_rows:
            for _ in range(N - current_reserved_ts_rows):
                t.Rows.Add(t.Rows(t.Rows.Count))
        elif N < current_reserved_ts_rows:
            for _ in range(current_reserved_ts_rows - N):
                t.Rows(t.Rows.Count - 1).Delete()
        
        sums = {"mech": 0.0, "prog": 0.0, "policy": 0.0, "effect": 0.0, "total": 0.0}
        count_evaluated = 0
        
        # 如果未传入乡镇与抽样村对应关系，尝试构建映射
        ts_v_map = township_villages_map or {}
        
        for idx, (code, name) in enumerate(township_list):
            r_idx = idx + 3
            t.Rows(r_idx).Cells(1).Range.Text = str(idx + 2)
            t.Rows(r_idx).Cells(2).Range.Text = name
            
            # 获取该乡镇本级的机制运行得分 (qsdwdm == code 或 level == 'township')
            township_rec = records_by_qsdwdm.get(code)
            township_self_mech = None
            if township_rec:
                ts_scores = calculate_neiye_subscores(township_rec.get("form_data", {}))["score"]
                township_self_mech = ts_scores["mech"]

            # 获取该乡镇下所有抽样村在 neiye_records 中的记录
            v_codes = ts_v_map.get(name, [])
            village_recs = []
            if v_codes:
                for vc in v_codes:
                    if vc in records_by_qsdwdm:
                        village_recs.append(records_by_qsdwdm[vc])
            else:
                # 兼容：前缀以该镇代码开头且长度>9位、或等于该镇代码的记录
                for r_code, r_data in records_by_qsdwdm.items():
                    if r_code != county_code and (r_code == code or (r_code.startswith(code) and len(r_code) > 9)):
                        village_recs.append(r_data)
            
            if village_recs or township_self_mech is not None:
                count_evaluated += 1
                v_scores = [calculate_neiye_subscores(vr.get("form_data", {}))["score"] for vr in village_recs]
                k = len(v_scores)
                
                # 计算下属抽样村各项指标平均分
                if k > 0:
                    village_avg_mech = sum(s["mech"] for s in v_scores) / k
                    sc_prog = sum(s["prog"] for s in v_scores) / k
                    sc_policy = sum(s["policy"] for s in v_scores) / k
                    sc_effect = sum(s["effect"] for s in v_scores) / k
                else:
                    village_avg_mech = None
                    sc_prog = 0.0
                    sc_policy = 0.0
                    sc_effect = 0.0

                # 机制运行得分计算方式：该乡镇机制运行得分与乡镇下属村的机制运行平均分
                if township_self_mech is not None and village_avg_mech is not None:
                    sc_mech = (township_self_mech + village_avg_mech) / 2.0
                elif township_self_mech is not None:
                    sc_mech = township_self_mech
                elif village_avg_mech is not None:
                    sc_mech = village_avg_mech
                else:
                    sc_mech = 15.0

                sc_total = sc_mech + sc_prog + sc_policy + sc_effect

                t.Rows(r_idx).Cells(3).Range.Text = f"{sc_mech:.1f}".rstrip('0').rstrip('.')
                t.Rows(r_idx).Cells(4).Range.Text = f"{sc_prog:.1f}".rstrip('0').rstrip('.') if k > 0 else "/"
                t.Rows(r_idx).Cells(5).Range.Text = f"{sc_policy:.1f}".rstrip('0').rstrip('.') if k > 0 else "/"
                t.Rows(r_idx).Cells(6).Range.Text = f"{sc_effect:.1f}".rstrip('0').rstrip('.') if k > 0 else "/"
                t.Rows(r_idx).Cells(7).Range.Text = f"{sc_total:.1f}".rstrip('0').rstrip('.')
                
                sums["mech"] += sc_mech
                sums["prog"] += sc_prog
                sums["policy"] += sc_policy
                sums["effect"] += sc_effect
                sums["total"] += sc_total
            else:
                t.Rows(r_idx).Cells(3).Range.Text = ""
                t.Rows(r_idx).Cells(4).Range.Text = ""
                t.Rows(r_idx).Cells(5).Range.Text = ""
                t.Rows(r_idx).Cells(6).Range.Text = ""
                t.Rows(r_idx).Cells(7).Range.Text = ""

        # Last Row: 总评 (平均分)
        summary_row = t.Rows.Count
        t.Rows(summary_row).Cells(1).Range.Text = str(N + 2)
        t.Rows(summary_row).Cells(2).Range.Text = "总评"
        if count_evaluated > 0:
            mech_count = count_evaluated + 1
            avg_mech = (sums['mech'] + county_mech) / mech_count
            avg_prog = sums['prog'] / count_evaluated
            avg_policy = sums['policy'] / count_evaluated
            avg_effect = sums['effect'] / count_evaluated
            avg_total = avg_mech + avg_prog + avg_policy + avg_effect
            
            t.Rows(summary_row).Cells(3).Range.Text = f"{avg_mech:.1f}"
            t.Rows(summary_row).Cells(4).Range.Text = f"{avg_prog:.1f}"
            t.Rows(summary_row).Cells(5).Range.Text = f"{avg_policy:.1f}"
            t.Rows(summary_row).Cells(6).Range.Text = f"{avg_effect:.1f}"
            t.Rows(summary_row).Cells(7).Range.Text = f"{avg_total:.1f}"
        else:
            t.Rows(summary_row).Cells(3).Range.Text = ""
            t.Rows(summary_row).Cells(4).Range.Text = ""
            t.Rows(summary_row).Cells(5).Range.Text = ""
            t.Rows(summary_row).Cells(6).Range.Text = ""
            t.Rows(summary_row).Cells(7).Range.Text = ""

        doc.Save()
        doc.Close(False)
        word.Quit()
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_neiye_att7 error:", e)
        try: word.Quit()
        except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()
def export_att4(township_name, farmer_count, total_area):
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        
        base_dir = get_base_dir()
        os.makedirs(os.path.join(base_dir, 'backend', 'downloads'), exist_ok=True)
        template = os.path.join(base_dir, '附件', '附件4.doc')
        clean_ts = sanitize_filename(township_name)
        out = os.path.join(base_dir, 'backend', 'downloads', f'附件4_成果检查验收申请表_{clean_ts}.doc')
        if os.path.exists(out):
            try: os.remove(out)
            except: pass
        shutil.copy(template, out)
        
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")

        doc = word.Documents.Open(FileName=out, ReadOnly=False, ConfirmConversions=False)
        replace_common_bookmarks(doc, county_name)
        t = doc.Tables(1)
        t.Cell(1, 2).Range.Text = township_name
        t.Cell(5, 2).Range.Text = str(farmer_count)
        t.Cell(6, 2).Range.Text = f"{total_area:.2f}"
        
        doc.SaveAs2(FileName=out, FileFormat=0)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/附件4_成果检查验收申请表_{clean_ts}.doc"
    except Exception as e:
        print("Export att4 error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_att5(stats_data, township_code, township_name):
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        
        base_dir = get_base_dir()
        os.makedirs(os.path.join(base_dir, 'backend', 'downloads'), exist_ok=True)
        
        att5_template = os.path.join(base_dir, '附件', '附件5.doc')
        clean_ts = sanitize_filename(township_name)
        att5_out = os.path.join(base_dir, 'backend', 'downloads', f'附件5_抽样统计表_{clean_ts}.doc')
        if os.path.exists(att5_out):
            try: os.remove(att5_out)
            except: pass
        shutil.copy(att5_template, att5_out)
        
        doc5 = word.Documents.Open(FileName=att5_out, ReadOnly=False, ConfirmConversions=False)
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        replace_common_bookmarks(doc5, county_name)
        t5 = doc5.Tables(1)
        
        N = len(stats_data)
        while t5.Rows.Count > 2:
            t5.Rows(3).Delete()
            
        if N > 1:
            t5.Rows(2).Select()
            word.Selection.InsertRowsBelow(N - 1)
            
        for i, row in enumerate(stats_data):
            r_idx = i + 2
            try:
                t5.Cell(r_idx, 1).Range.Text = str(row.get('序号', ''))
                t5.Cell(r_idx, 2).Range.Text = str(row.get('乡镇名称', ''))
                t5.Cell(r_idx, 3).Range.Text = str(row.get('村名称', ''))
                t5.Cell(r_idx, 4).Range.Text = str(row.get('组名称', ''))
                t5.Cell(r_idx, 5).Range.Text = str(row.get('发包方总户数', ''))
                t5.Cell(r_idx, 6).Range.Text = str(row.get('抽样农户数5%', ''))
            except: pass
            
        doc5.SaveAs2(FileName=att5_out, FileFormat=0)
        doc5.Close(0)
        doc5 = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/附件5_抽样统计表_{clean_ts}.doc"
    except Exception as e:
        print("Export att5 error:", e)
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_docs(stats_data, att8_data, township_code, township_name):
    url = export_att5(stats_data, township_code, township_name)
    return [url]

def export_waiye_att8(township_name, village_name, group_name, group_rows):
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        word.ScreenUpdating = False
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, '附件', '附件8.doc')
        clean_ts = sanitize_filename(township_name)
        clean_vn = sanitize_filename(village_name)
        clean_gn = sanitize_filename(group_name)
        
        os.makedirs(os.path.join(base_dir, 'backend', 'downloads'), exist_ok=True)
        out_path = os.path.join(base_dir, 'backend', 'downloads', f'附件8_外业核查记录表_{clean_ts}{clean_vn}{clean_gn}.doc')
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass
        shutil.copy(tpl, out_path)
        
        doc8 = word.Documents.Open(FileName=out_path, ReadOnly=False, ConfirmConversions=False)
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        replace_common_bookmarks(doc8, county_name)

        now = datetime.now()
        date_str = f"{now.year} 年  {now.month:02d}  月  {now.day:02d}  日"
        p3 = doc8.Paragraphs(3)
        rng = p3.Range
        rng.End = rng.End - 1
        rng.Text = f"   乡镇：{township_name} \t行政村：{village_name} \t村民小组：{group_name}" + " "*25 + date_str
        
        t8 = doc8.Tables(1)
        for _ in range(5):
            try: t8.Rows(3).Delete()
            except: pass
            
        N_rows = len(group_rows)
        if N_rows > 1:
            t8.Rows(2).Select()
            word.Selection.InsertRowsBelow(N_rows - 1)
            
        total_errors = 0
        satisfaction_count = 0
        
        # Group by contractor to calculate unique contractor-level errors and satisfaction
        processed_contractors = set()
        contractor_count = 0

        for i, r in enumerate(group_rows):
            r_idx = i + 2
            farmer_name = r.get('cbfmc', '') or r.get('承包方代表', '')
            cbfbm = str(r.get("cbfbm", "") or r.get("cbfbm_short", "") or r.get("cbfmc", ""))
            is_first_parcel_of_cbf = (cbfbm not in processed_contractors)
            
            # Parcel-level errors
            for k in ['area_acknowledged', 'bound_correct', 'self_verified']:
                if r.get(k) == 'X':
                    total_errors += 1
            
            lxdh_val = str(r.get('lxdh', '') or r.get('联系电话', '') or '').strip()
            if lxdh_val == 'None':
                lxdh_val = ''
            is_phone_err = (str(r.get('phone_correct', '')).strip() == 'X')

            # Contractor-level errors & satisfaction
            sat = r.get('satisfaction', '满意')
            if is_first_parcel_of_cbf:
                processed_contractors.add(cbfbm)
                contractor_count += 1
                for k in ['rights_correct', 'member_qualified', 'self_signed']:
                    if r.get(k) == 'X':
                        total_errors += 1
                if is_phone_err:
                    total_errors += 1
                if sat == '满意':
                    satisfaction_count += 1

            # 序号、地块信息每行必填
            t8.Cell(r_idx, 1).Range.Text = str(i + 1)
            t8.Cell(r_idx, 5).Range.Text = str(r.get('dkmc', '') or r.get('地块名称', ''))
            t8.Cell(r_idx, 6).Range.Text = str(r.get('dkbm_short', '') or r.get('地块简编码', ''))
            t8.Cell(r_idx, 7).Range.Text = str(r.get('scmj', '') or r.get('成果面积(亩)', ''))

            # 地块级核查指标
            for c_pos, k_name in [(8, 'area_acknowledged'), (10, 'bound_correct'), (12, 'self_verified')]:
                cell_k = t8.Cell(r_idx, c_pos)
                val_k = r.get(k_name, '')
                if val_k == 'X':
                    cell_k.Range.Text = 'X'
                    cell_k.Range.Font.Color = 255
                else:
                    cell_k.Range.Text = '√'

            # 农户级指标（仅首行填值并设颜色，后续行保持空值，便于后续单元格纵向合并且不重复文字）
            if is_first_parcel_of_cbf:
                t8.Cell(r_idx, 2).Range.Text = farmer_name
                t8.Cell(r_idx, 3).Range.Text = str(r.get('cbfbm_short', '') or r.get('承包方编码(缩略码)', ''))
                
                cell_dh = t8.Cell(r_idx, 4)
                if is_phone_err:
                    cell_dh.Range.Text = 'X'
                    cell_dh.Range.Font.Color = 255
                else:
                    cell_dh.Range.Text = lxdh_val if lxdh_val else '/'

                for c_pos, k_name in [(9, 'rights_correct'), (11, 'member_qualified'), (13, 'self_signed')]:
                    cell_k = t8.Cell(r_idx, c_pos)
                    val_k = r.get(k_name, '')
                    if val_k == 'X':
                        cell_k.Range.Text = 'X'
                        cell_k.Range.Font.Color = 255
                    else:
                        cell_k.Range.Text = '√'

                t8.Cell(r_idx, 14).Range.Text = sat
                t8.Cell(r_idx, 15).Range.Text = r.get('survey_method', r.get('调查抽样方式', '现场'))
                t8.Cell(r_idx, 16).Range.Text = ""

        # Fill stats rows BEFORE any vertical merges
        total_count = len(group_rows)
        prog_score = max(20.0 - total_errors * 0.5, 0.0)
        effect_score = (satisfaction_count / contractor_count * 10.0) if contractor_count > 0 else 10.0
        
        r_stat_idx = total_count + 2
        t8.Cell(r_stat_idx, 2).Range.Text = str(total_errors)
        t8.Cell(r_stat_idx, 3).Range.Text = str(satisfaction_count)
        
        r_score_idx = total_count + 3
        t8.Cell(r_score_idx, 1).Range.Text = f"发包方程序规范得分=20-发包方错误总和({total_errors})×0.5={prog_score:.1f}；发包方工作成效（满意度调查）得分=满意数({satisfaction_count})/抽检数({contractor_count})×10={effect_score:.1f}"

        # Group contiguous rows by contractor (cbfbm) and merge Column 16 from bottom to top
        segments = []
        curr_cbfbm = None
        start_r = 2
        for i, r in enumerate(group_rows):
            r_idx = i + 2
            cbfbm = str(r.get("cbfbm", "") or r.get("cbfbm_short", "") or r.get("cbfmc", ""))
            if cbfbm != curr_cbfbm:
                if curr_cbfbm is not None:
                    segments.append((curr_cbfbm, start_r, r_idx - 1))
                curr_cbfbm = cbfbm
                start_r = r_idx
        if curr_cbfbm is not None:
            segments.append((curr_cbfbm, start_r, len(group_rows) + 1))

        sig_dir = os.path.join(base_dir, "backend", "uploads", "signatures")
        
        # 保持 100% 完整原版合并逻辑：合并全部 9 个农户级列
        for cbfbm, r_start, r_end in reversed(segments):
            if r_start < r_end:
                for col_idx in [16, 15, 14, 13, 11, 9, 4, 3, 2]:
                    try:
                        t8.Cell(r_start, col_idx).Merge(t8.Cell(r_end, col_idx))
                    except Exception as me:
                        print(f"Merge err at rows {r_start}-{r_end} col {col_idx}: {me}")
                cell_target = t8.Cell(r_start, 16)
            else:
                cell_target = t8.Cell(r_start, 16)

            sig_path = os.path.join(sig_dir, f"{cbfbm}.png")
            if os.path.exists(sig_path):
                cell_target.Range.Text = ""
                pic = cell_target.Range.InlineShapes.AddPicture(
                    FileName=os.path.abspath(sig_path), LinkToFile=False, SaveWithDocument=True
                )
                pic.Width = 65
                pic.Height = 26
                cell_target.Range.ParagraphFormat.Alignment = 1
                cell_target.VerticalAlignment = 1
            else:
                cell_target.Range.Text = ""
                
            sig_path = os.path.join(sig_dir, f"{cbfbm}.png")
            if os.path.exists(sig_path):
                cell_target.Range.Text = ""
                pic = cell_target.Range.InlineShapes.AddPicture(
                    FileName=os.path.abspath(sig_path), LinkToFile=False, SaveWithDocument=True
                )
                pic.Width = 65
                pic.Height = 26
                cell_target.Range.ParagraphFormat.Alignment = 1
                cell_target.VerticalAlignment = 1
            else:
                cell_target.Range.Text = ""

        try:
            t8.Range.ParagraphFormat.Alignment = 1
        except:
            pass
        doc8.SaveAs2(FileName=out_path, FileFormat=0)
        doc8.Close(0)
        doc8 = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/附件8_外业核查记录表_{clean_ts}{clean_vn}{clean_gn}.doc"
    except Exception as e:
        print("Export waiye att8 error:", e)
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_waiye_att9(samples_rows):
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        from database import get_county_and_townships_sync
        county_info, township_list_raw = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        all_townships = [item["name"] for item in township_list_raw]

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件9.doc")
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        out_filename = f"附件9_{county_name}县级自查外业组检查得分表.doc"
        out_path = os.path.join(base_dir, "backend", "downloads", out_filename)
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass
        shutil.copy(tpl, out_path)
        
        doc = word.Documents.Open(FileName=out_path, ReadOnly=False, ConfirmConversions=False)
        replace_common_bookmarks(doc, county_name)
        t = doc.Tables(1)
        
        # 1. 计算每个组的得分，并按 乡镇 -> 行政村 -> 村民小组 组织层级
        groups_map = defaultdict(list)
        for r in samples_rows:
            key = (r.get("township_name", ""), r.get("village_name", ""), r.get("group_name", ""))
            if key[0]:
                groups_map[key].append(r)
            
        # 结构：hierarchy[t_name][v_name] = [ { "group": g_name, "prog_score": ..., "effect_score": ... }, ... ]
        hierarchy = defaultdict(lambda: defaultdict(list))
        
        for (t_name, v_name, g_name), g_rows in groups_map.items():
            total_errors = 0
            satisfaction_count = 0
            processed_cbfs = set()
            c_cnt = 0
            for r in g_rows:
                for k in ["area_acknowledged", "bound_correct", "self_verified"]:
                    if r.get(k) == "X":
                        total_errors += 1
                
                cbf_key = str(r.get("cbfbm") or r.get("cbfbm_short") or r.get("cbfmc") or "")
                if cbf_key not in processed_cbfs:
                    processed_cbfs.add(cbf_key)
                    c_cnt += 1
                    for k in ["rights_correct", "member_qualified", "self_signed"]:
                        if r.get(k) == "X":
                            total_errors += 1
                    if str(r.get('phone_correct', '')).strip() == 'X':
                        total_errors += 1
                    if r.get("satisfaction") == "满意":
                        satisfaction_count += 1
                    
            c_cnt = c_cnt if c_cnt > 0 else len(g_rows)
            prog_score = max(20.0 - total_errors * 0.5, 0.0)
            effect_score = (satisfaction_count / c_cnt * 10.0) if c_cnt > 0 else 10.0
            
            hierarchy[t_name][v_name].append({
                "group": g_name,
                "prog_score": prog_score,
                "effect_score": effect_score
            })
            
        # 2. 按照实际有抽样数据的乡镇排序列出块结构
        # 规则：村得分由各组得分平均值，乡镇得分为各村平均值
        if all_townships:
            sampled_townships = [t for t in all_townships if t in hierarchy]
        else:
            sampled_townships = list(hierarchy.keys())

        township_blocks = []
        current_row = 3
        total_rows_needed = 0

        for ts_idx, t_name in enumerate(sampled_townships):
            v_dict = hierarchy.get(t_name, {})
            if not v_dict:
                continue

            ts_start_row = current_row
            v_blocks = []
            
            for v_name, g_list in v_dict.items():
                if not g_list:
                    continue
                v_start_row = current_row
                g_count = len(g_list)
                v_end_row = current_row + g_count - 1
                
                # 村得分为各组得分平均值
                v_prog = sum(g["prog_score"] for g in g_list) / g_count
                v_effect = sum(g["effect_score"] for g in g_list) / g_count
                
                v_blocks.append({
                    "village": v_name,
                    "v_prog": v_prog,
                    "v_effect": v_effect,
                    "start_row": v_start_row,
                    "end_row": v_end_row,
                    "groups": g_list
                })
                current_row += g_count
                total_rows_needed += g_count

            ts_end_row = current_row - 1
            # 乡镇得分为各村得分平均值
            if v_blocks:
                town_prog = sum(vb["v_prog"] for vb in v_blocks) / len(v_blocks)
                town_effect = sum(vb["v_effect"] for vb in v_blocks) / len(v_blocks)
            else:
                town_prog = 20.0
                town_effect = 10.0

            township_blocks.append({
                "ts_no": str(ts_idx + 1),
                "township": t_name,
                "town_prog": town_prog,
                "town_effect": town_effect,
                "start_row": ts_start_row,
                "end_row": ts_end_row,
                "villages": v_blocks
            })

        N = total_rows_needed
        # 表格初始结构：Row 1 和 2 是表头，预留数据行数 = t.Rows.Count - 2
        current_reserved = t.Rows.Count - 2
        if N > current_reserved:
            last_row = t.Rows.Count
            t.Cell(last_row, 1).Select()
            word.Selection.InsertRowsBelow(N - current_reserved)
        elif N < current_reserved:
            for _ in range(current_reserved - N):
                last_row = t.Rows.Count
                t.Cell(last_row, 1).Select()
                word.Selection.Rows.Delete()

        # 3. 填充全部单元格
        # 对于即将纵向合并的列：
        # - 镇级（列 1~4）：仅在 ts_start_row 填充文本，其余行留空
        # - 村级（列 5~7）：仅在 v_start_row 填充文本，其余行留空
        # - 组级（列 8~10）：每行填充对应的组信息与得分
        for ts_b in township_blocks:
            for vb in ts_b["villages"]:
                for g_idx, g_item in enumerate(vb["groups"]):
                    r_idx = vb["start_row"] + g_idx
                    
                    # 镇级列 (1~4)
                    if r_idx == ts_b["start_row"]:
                        t.Cell(r_idx, 1).Range.Text = ts_b["ts_no"]
                        t.Cell(r_idx, 2).Range.Text = ts_b["township"]
                        t.Cell(r_idx, 3).Range.Text = format_score(ts_b["town_prog"])
                        t.Cell(r_idx, 4).Range.Text = format_score(ts_b["town_effect"])
                    else:
                        t.Cell(r_idx, 1).Range.Text = ""
                        t.Cell(r_idx, 2).Range.Text = ""
                        t.Cell(r_idx, 3).Range.Text = ""
                        t.Cell(r_idx, 4).Range.Text = ""

                    # 村级列 (5~7)
                    if r_idx == vb["start_row"]:
                        t.Cell(r_idx, 5).Range.Text = vb["village"]
                        t.Cell(r_idx, 6).Range.Text = format_score(vb["v_prog"])
                        t.Cell(r_idx, 7).Range.Text = format_score(vb["v_effect"])
                    else:
                        t.Cell(r_idx, 5).Range.Text = ""
                        t.Cell(r_idx, 6).Range.Text = ""
                        t.Cell(r_idx, 7).Range.Text = ""

                    # 组级列 (8~10)
                    t.Cell(r_idx, 8).Range.Text = g_item["group"]
                    t.Cell(r_idx, 9).Range.Text = format_score(g_item["prog_score"])
                    t.Cell(r_idx, 10).Range.Text = format_score(g_item["effect_score"])

        # 4. 逆序执行单元格纵向合并 (先合并村级列 7,6,5，再合并镇级列 4,3,2,1)
        # 倒序遍历村块执行合并
        all_village_blocks = []
        for ts_b in township_blocks:
            for vb in ts_b["villages"]:
                all_village_blocks.append(vb)

        for vb in reversed(all_village_blocks):
            if vb["end_row"] > vb["start_row"]:
                for c in [7, 6, 5]:
                    try:
                        t.Cell(vb["start_row"], c).Merge(t.Cell(vb["end_row"], c))
                        t.Cell(vb["start_row"], c).VerticalAlignment = 1
                        t.Cell(vb["start_row"], c).Range.ParagraphFormat.Alignment = 1
                    except Exception as me:
                        print(f"Merge village cells error at {vb['start_row']}-{vb['end_row']} col {c}: {me}")

        # 倒序遍历乡镇块执行合并
        for ts_b in reversed(township_blocks):
            if ts_b["end_row"] > ts_b["start_row"]:
                for c in [4, 3, 2, 1]:
                    try:
                        t.Cell(ts_b["start_row"], c).Merge(t.Cell(ts_b["end_row"], c))
                        t.Cell(ts_b["start_row"], c).VerticalAlignment = 1
                        t.Cell(ts_b["start_row"], c).Range.ParagraphFormat.Alignment = 1
                    except Exception as me:
                        print(f"Merge township cells error at {ts_b['start_row']}-{ts_b['end_row']} col {c}: {me}")

        doc.SaveAs2(FileName=out_path, FileFormat=0)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_waiye_att9 error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

# ================= 自查整改（附件12 / 附件13） =================

WAIYE_FLAG_LABELS = [
    ("area_acknowledged", "面积不认可"),
    ("rights_correct", "权属不正确"),
    ("bound_correct", "四至不正确"),
    ("member_qualified", "成员不符合"),
    ("self_verified", "非本人核实"),
    ("self_signed", "非本人签名"),
    ("phone_correct", "联系电话不正确"),
]

def _build_rectify_rows(neiye_form, waiye_rows, township_name):
    rows = []
    
    MAJOR_NEIYE = [
        "未制定方案", "直接套用上级方案", "制定程序不合法", 
        "未召开会议", "未公示工作组名单", "参会人数不足法定数量", 
        "没有延包方案", "延包方案未上报", "延包方案未公示", "未召开会议讨论延包方案", 
        "合同版本格式不正确", "合同网签率未达到95%", "没有地块示意图", 
        "档案整理第三方无涉密档案整理资质", "档案资料收集不齐全", "档案分类不符合要求", 
        "档号与归档章不一致", "档案变数字化瑕疵", "姓名著录错误"
    ]
    SEVERE_NEIYE = [
        "支付不规范", "支付不及时", "经费没有县级兜底", "没有进行摸底", 
        "摸底表农户未签署", "摸底表中没有表达延包意愿", "摸底表其它签署不齐全", 
        "特殊人员摸底不清或未统计", "户变化未统计", "矛盾纠纷未登记或处理不当", 
        "承包地变化未摸清", "没有应确尽确", "没有公示材料", "没有公示不足15天", "没有公示或不足15天", "公示不足15天",
        "公示结果未确认", "未进行信息共享", "未与不动产登记部门有序衔接", 
        "未保障特殊群体权益", "未保障无地户权益", "没有应收尽收", 
        "采用不正当方式隐匿消亡户", "违背农户意愿强行推进", "确权确股不确地手续不齐全", 
        "小调整比率过大或手续不齐全", "打乱重分", "违法调整或收回承包地", 
        "未建立矛盾纠纷处置机制", "未建立舆情处置办法", "没有矛盾纠纷处理台账"
    ]
    
    MAJOR_WAIYE = ["权属不正确", "成员不符合"]
    SEVERE_WAIYE = ["非本人签名", "面积不认可"]
    
    def get_level(problem_desc, is_neiye=True):
        if is_neiye:
            if any(m in problem_desc for m in MAJOR_NEIYE):
                return "重大"
            if any(s in problem_desc for s in SEVERE_NEIYE):
                return "较重"
            return "一般"
        else:
            if any(m in problem_desc for m in MAJOR_WAIYE):
                return "重大"
            if any(s in problem_desc for s in SEVERE_WAIYE):
                return "较重"
            return "一般"

    # 内业问题：先归一化，兼容旧数据中字段为字符串的情况
    def _norm_list(v):
        if isinstance(v, str): return [v] if v else []
        if isinstance(v, list): return v
        return []
    
    if isinstance(neiye_form, dict):
        # 列表型问题字段
        for key, val in neiye_form.items():
            items = _norm_list(val)
            for item in items:
                if item:
                    item_str = str(item)
                    rows.append({
                        "source": "内业",
                        "level": get_level(item_str, is_neiye=True),
                        "desc": f"{township_name} {item_str}",
                        "scope": "",
                        "unit": township_name,
                    })
        # 数量型问题字段：未保障特殊/无地户权益、消亡户应收尽收、不正当隐匿消亡户
        policy_fields = [
            ("policy_2_1", "未保障特殊群体权益"),
            ("policy_2_2", "未保障无地户权益"),
            ("policy_3_1", "没有应收尽收"),
            ("policy_3_2", "采用不正当方式隐匿消亡户"),
        ]
        for key, label in policy_fields:
            cnt = int(neiye_form.get(key, 0) or 0)
            if cnt > 0:
                rows.append({
                    "source": "内业",
                    "level": get_level(label, is_neiye=True),
                    "desc": f"{township_name} {label}（{cnt}起）",
                    "scope": str(cnt),
                    "unit": township_name,
                })
    # 外业问题：waiye_samples 中按户和地块分组
    CBF_ERRORS = [
        ("phone_correct", "联系电话不正确"),
        ("member_qualified", "成员不符合"),
        ("rights_correct", "权属不正确"),
        ("self_signed", "非本人签名"),
    ]
    DK_ERRORS = [
        ("bound_correct", "四至不正确"),
        ("area_acknowledged", "面积不认可"),
        ("self_verified", "非本人核实"),
    ]

    from collections import defaultdict
    waiye_groups = defaultdict(list)
    for r in (waiye_rows or []):
        cbf_key = (r.get("village_name", ""), r.get("group_name", ""), r.get("cbfmc", ""), r.get("cbfbm_short", ""))
        waiye_groups[cbf_key].append(r)

    for (village_name, group_name, cbfmc, cbfbm_short), group_rows in waiye_groups.items():
        first_r = group_rows[0]
        cbf_failed = [lbl for key, lbl in CBF_ERRORS if str(first_r.get(key, "")).strip() == "X"]
        
        dk_failed_list = []
        for r in group_rows:
            dk_failed = [lbl for key, lbl in DK_ERRORS if str(r.get(key, "")).strip() == "X"]
            if dk_failed:
                dk_name = r.get("dkbm_short", "") or r.get("dkmc", "")
                dk_failed_list.append(f"{dk_name}地块，{'、'.join(dk_failed)}")
                
        if cbf_failed or dk_failed_list:
            cbf_suffix = f"({cbfbm_short})" if cbfbm_short else ""
            loc = f"{township_name}{village_name}{group_name} {cbfmc}户{cbf_suffix}"
            desc_parts = []
            if cbf_failed:
                desc_parts.append(f"承包方核查问题有：{'，'.join(cbf_failed)}。")
            if dk_failed_list:
                desc_parts.append(f"地块核查问题有：{'；'.join(dk_failed_list)}。")
                
            desc_str = "\n".join(desc_parts)
            lvl = get_level(desc_str, is_neiye=False)
            
            rows.append({
                "source": "外业",
                "level": lvl,
                "desc": f"{loc}，{desc_str}",
                "scope": "",
                "unit": village_name or township_name,
            })
            
    if not rows:
        rows.append({
            "source": "",
            "level": "",
            "desc": "经自查，暂未发现需整改的问题",
            "scope": "",
            "unit": township_name,
        })
    return rows

def export_rectify_att12(township_name):
    pythoncom.CoInitialize()
    word = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件12.doc")
        clean_ts = sanitize_filename(township_name)
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        out_path = os.path.join(base_dir, "backend", "downloads", f"附件12_整改通知书_{clean_ts}.doc")
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass
        shutil.copy(tpl, out_path)

        doc = word.Documents.Open(FileName=out_path, ReadOnly=False, ConfirmConversions=False)

        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        replace_common_bookmarks(doc, county_name)

        # 替换标题中的"（模版）"，用 Find 保留段落格式
        f12 = doc.Content.Find
        f12.ClearFormatting()
        f12.Replacement.ClearFormatting()
        f12.Execute(FindText="（模版）", ReplaceWith="", Replace=2)
        f12.Execute(FindText="（模板）", ReplaceWith="", Replace=2)

        # 替换书签 XJQYMC 为乡镇名称
        if doc.Bookmarks.Exists("XJQYMC"):
            doc.Bookmarks("XJQYMC").Range.Text = township_name

        doc.SaveAs2(FileName=out_path, FileFormat=0)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/附件12_整改通知书_{clean_ts}.doc"
    except Exception as e:
        print("export_rectify_att12 error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_rectify_att13(township_name, neiye_form, waiye_rows):
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "附件13.doc")
        clean_ts = sanitize_filename(township_name)
        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        out_path = os.path.join(base_dir, "backend", "downloads", f"附件13_问题整改销号台账_{clean_ts}.doc")
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass
        shutil.copy(tpl, out_path)

        doc = word.Documents.Open(FileName=out_path, ReadOnly=False, ConfirmConversions=False)
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        replace_common_bookmarks(doc, county_name)

        t = doc.Tables(1)

        # 删除模板中的空白数据行，仅保留表头
        while t.Rows.Count > 1:
            t.Rows(t.Rows.Count).Delete()

        data_rows = _build_rectify_rows(neiye_form, waiye_rows, township_name)

        for i, item in enumerate(data_rows):
            # 始终选中当前最后一行，在其下方插入新行，保证数据顺序正确
            t.Rows(t.Rows.Count).Select()
            word.Selection.InsertRowsBelow(1)
            r = i + 2
            
            # 取消新建行的表头属性（防止从第一行继承导致全表变成表头无法跨页重复）
            t.Rows(r).HeadingFormat = False
            
            row_rng = t.Rows(r).Range
            row_rng.ParagraphFormat.Alignment = 1    # 居中
            row_rng.Font.Bold = False                # 取消加粗
            for c in range(1, 17):
                t.Cell(r, c).Range.Text = ""
            t.Cell(r, 1).Range.Text = str(i + 1)                       # 序号
            t.Cell(r, 2).Range.Text = item["source"]                   # 问题来源
            t.Cell(r, 3).Range.Text = item.get("level", "一般")        # 问题等级
            t.Cell(r, 4).Range.Text = item["desc"]                     # 问题具体描述
            t.Cell(r, 5).Range.Text = item.get("scope", "")            # 涉及资料/农户数量
            t.Cell(r, 6).Range.Text = item.get("unit", "")             # 整改责任单位
            t.Cell(r, 9).Range.Text = "对照问题逐项整改，补齐材料，规范程序"  # 整改措施
            # 列13 复核情况、列15 销号状态保持为空

        # 在所有数据行插入完毕后，重新设置表头重复标题行（确保生效）
        t.Rows(1).HeadingFormat = True

        doc.SaveAs2(FileName=out_path, FileFormat=0)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/附件13_问题整改销号台账_{clean_ts}.doc"
    except Exception as e:
        print("export_rectify_att13 error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()

def export_waiye_inquiry(data):
    import os, shutil, win32com.client, pythoncom
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        
        base_dir = get_base_dir()
        tpl = os.path.join(base_dir, "附件", "询问笔录.doc")
        
        # 实际承包方名称用于文件名和书签 cbfmc
        real_cbfmc = str(data.get("cbfmc", "")).strip()
        clean_cbf = sanitize_filename(real_cbfmc)

        # 文件名前缀带上 XX镇XX村
        clean_ts = sanitize_filename(data.get("township_name", "") or "")
        clean_vn = sanitize_filename(data.get("village_name", "") or "")
        loc_parts = [p for p in [clean_ts, clean_vn] if p]
        loc_prefix = "".join(loc_parts)
        if loc_prefix:
            out_filename = f"附件_询问笔录_{loc_prefix}_{clean_cbf}.doc"
        else:
            out_filename = f"附件_询问笔录_{clean_cbf}.doc"

        os.makedirs(os.path.join(base_dir, "backend", "downloads"), exist_ok=True)
        out_path = os.path.join(base_dir, "backend", "downloads", out_filename)
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass
        shutil.copy(tpl, out_path)
        
        doc = word.Documents.Open(FileName=out_path, ReadOnly=False, ConfirmConversions=False)
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")
        replace_common_bookmarks(doc, county_name)

        fd = data.get("form_data", {})
        
        def _check(doc, target_text):
            f = doc.Content.Find
            f.ClearFormatting()
            f.Replacement.ClearFormatting()
            f.Execute("□" + target_text, False, False, False, False, False, True, 1, False, "☑" + target_text, 2)
            
        def _fill_bm(doc, bm_name, text, underline=False):
            if doc.Bookmarks.Exists(bm_name):
                rng = doc.Bookmarks(bm_name).Range
                if underline:
                    # Pad text with spaces for better visual underline
                    text = f" {text} "
                rng.Text = text
                if underline:
                    rng.Font.Underline = 1
                doc.Bookmarks.Add(Name=bm_name, Range=rng)
                
        import datetime
        now = datetime.datetime.now()
        
        _fill_bm(doc, "year", str(now.year), True)
        _fill_bm(doc, "month", str(now.month).zfill(2), True)
        _fill_bm(doc, "day", str(now.day).zfill(2), True)
        _fill_bm(doc, "hour", str(now.hour).zfill(2), True)
        _fill_bm(doc, "minite", str(now.minute).zfill(2), True)
        
        _fill_bm(doc, "xjqymc", data.get('township_name',''), True)
        _fill_bm(doc, "cjqymc", data.get('village_name',''), True)
        _fill_bm(doc, "zjqymc", data.get('group_name',''), True)
        
        # 书签 bxwr -> 页面基本信息获取 (被询问人姓名)
        bxwr_name = data.get("bxwr") or fd.get("bxwr") or fd.get("cbfmc") or real_cbfmc
        _fill_bm(doc, "bxwr", bxwr_name, True)
        
        # 书签 cbfmc -> 实际承包方名称
        _fill_bm(doc, "cbfmc", real_cbfmc, False)
        
        _fill_bm(doc, "xb", data.get('gender','男'), True)
        _fill_bm(doc, "lxdh", data.get('lxdh',''), True)
        
        place = fd.get('inquiry_place', '')
        if place in ['农户家中', '田间地头', '村委会']:
            _check(doc, place)
                    
        rel = fd.get('relationship', '')
        if rel in ['本人', '配偶', '子女']:
            _check(doc, rel)
        elif rel == '其他亲属':
            _check(doc, "其他亲属")
            
        _fill_bm(doc, "xwr", fd.get('inquirer',''), True)
        

        questions = fd.get('questions', [])
        xwnr_text = ""
        
        idx = 1
        for q in questions:
            if not q.get('checked'):
                continue
                
            q_text = q.get('q', '')
            xwnr_text += f"{idx}、{q_text}\r"
            
            opts = q.get('opts', [])
            ans = q.get('answer', '')
            has_desc = q.get('has_desc', False)
            desc_label = q.get('desc_label', '')
            desc = q.get('desc', '')
            
            ans_line = "答："
            if opts:
                opt_strs = []
                for o in opts:
                    if o == ans:
                        opt_strs.append(f"☑{o}")
                    else:
                        opt_strs.append(f"□{o}")
                ans_line += " ".join(opt_strs)
                
            if has_desc:
                if desc_label == "答：":
                    if not opts:
                        ans_line += f"{desc}"
                    else:
                        ans_line += f"  {desc}"
                else:
                    ans_line += f"  {desc_label}{desc}"
                
            xwnr_text += ans_line + "\r"
            idx += 1
            
        if doc.Bookmarks.Exists("wxnr"):
                rng = doc.Bookmarks("wxnr").Range
                rng.Text = xwnr_text
                rng.Font.Size = 12  # 小四号
                rng.ParagraphFormat.LineSpacingRule = 1  # wdLineSpace1pt5
                rng.ParagraphFormat.SpaceBefore = 0
                rng.ParagraphFormat.SpaceAfter = 0
                doc.Bookmarks.Add(Name="wxnr", Range=rng)
        
        # Signatures (取消村民代表签名，仅保留被询问人与询问人签名)
        sig_bxwrqm = os.path.join(base_dir, "backend", "uploads", "signatures", f"{data.get('cbfbm')}_bxwrqm.png")
        sig_xwrqm = os.path.join(base_dir, "backend", "uploads", "signatures", f"{data.get('cbfbm')}_xwrqm.png")

        def _insert_sig(doc, bm_name, path):
            if os.path.exists(path):
                if doc.Bookmarks.Exists(bm_name):
                    rng_sig = doc.Bookmarks(bm_name).Range
                    # Collapse to start
                    rng_sig.Collapse(1)
                    pic = rng_sig.InlineShapes.AddPicture(FileName=os.path.abspath(path), LinkToFile=False, SaveWithDocument=True)
                    pic.Width = 65
                    pic.Height = 26

        _insert_sig(doc, "bxwrqm", sig_bxwrqm)
        _insert_sig(doc, "xwrqm", sig_xwrqm)
        
        # 附带现场照片：在笔录末尾自动分页插入现场问询照片及图题
        photos = data.get("photos") or fd.get("photos") or []
        valid_photos = []
        for p_item in photos:
            raw_url = p_item if isinstance(p_item, str) else (p_item.get("url") or "")
            if not raw_url:
                continue
            clean_url = raw_url.strip()
            # 兼容多种 URL 表达形式
            if clean_url.startswith("/uploads/"):
                rel_path = clean_url.lstrip("/")
            elif "file=" in clean_url:
                rel_path = clean_url.split("file=")[-1].lstrip("/")
            else:
                rel_path = os.path.join("uploads", os.path.basename(clean_url))

            # 兼容各种工作目录层级
            candidates = [
                os.path.join(base_dir, "backend", rel_path),
                os.path.join(base_dir, rel_path),
                os.path.join(os.path.dirname(__file__), rel_path),
                os.path.abspath(clean_url)
            ]
            for cand in candidates:
                if os.path.exists(cand) and os.path.isfile(cand):
                    valid_photos.append(cand)
                    break

        if valid_photos:
            try:
                # 在文档末尾插入分页符 (wdPageBreak = 7)
                rng_end = doc.Content
                rng_end.Collapse(0)  # wdCollapseEnd = 0
                rng_end.InsertBreak(7)

                # 插入附图大标题
                p_title = doc.Paragraphs.Add()
                p_title.Range.Text = "附：现场问询核查照片\r"
                p_title.Range.Font.Name = "黑体"
                p_title.Range.Font.Size = 14
                p_title.Range.Font.Bold = True
                p_title.Range.ParagraphFormat.Alignment = 1  # 居中
                p_title.Range.ParagraphFormat.SpaceBefore = 12
                p_title.Range.ParagraphFormat.SpaceAfter = 14

                for p_idx, photo_abs_path in enumerate(valid_photos, 1):
                    # 插入照片
                    p_img = doc.Paragraphs.Add()
                    shape = p_img.Range.InlineShapes.AddPicture(
                        FileName=photo_abs_path,
                        LinkToFile=False,
                        SaveWithDocument=True
                    )
                    # 适度缩放照片宽度至标准 350 pt (~12.3 cm)，高宽比锁定
                    shape.Width = 350
                    shape.Height = shape.Height * (350.0 / max(shape.Width, 1))
                    p_img.Range.ParagraphFormat.Alignment = 1  # 居中
                    p_img.Range.ParagraphFormat.SpaceAfter = 6

                    # 插入图题说明
                    p_cap = doc.Paragraphs.Add()
                    p_cap.Range.Text = f"现场照片 {p_idx}\r"
                    p_cap.Range.Font.Name = "宋体"
                    p_cap.Range.Font.Size = 10.5
                    p_cap.Range.Font.Bold = False
                    p_cap.Range.ParagraphFormat.Alignment = 1  # 居中
                    p_cap.Range.ParagraphFormat.SpaceAfter = 16
            except Exception as pe:
                print(f"插入现场问询照片异常 (已忽略): {pe}")

        doc.SaveAs2(FileName=out_path, FileFormat=0)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_waiye_inquiry error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()


def export_village_meeting_photos(township_name: str, village_name: str, photos: list) -> str:
    """
    导出指定行政村的现场会与问询现场照片文档 (.docx)
    模板文件：附件/现场会照片.docx
    命名格式：{township_name}{village_name}现场会照片.docx
    排版规范：
    - 使用模板中的大标题与基本信息，自动替换为当前乡镇、行政村名称及核查日期
    - 使用模板中的 2×2 表格，将照片逐张插入表格单元格中
    - 根据单元格大小与图片原始比例自动计算自适应缩放尺寸，居中显示且不破坏表格布局
    - 超过 4 张照片时自动新增表格行，无照片时友好提示
    """
    import os, shutil, win32com.client, pythoncom, datetime
    from PIL import Image, ImageOps
    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        from database import get_county_and_townships_sync
        county_info, _ = get_county_and_townships_sync()
        county_name = county_info.get("name", "全椒县")

        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0

        base_dir = get_base_dir()
        downloads_dir = os.path.join(base_dir, "backend", "downloads")
        os.makedirs(downloads_dir, exist_ok=True)

        clean_ts = sanitize_filename(township_name or "")
        clean_vn = sanitize_filename(village_name or "")
        loc_prefix = f"{clean_ts}{clean_vn}" if (clean_ts or clean_vn) else "现场会"
        out_filename = f"{loc_prefix}现场会照片.docx"
        out_path = os.path.join(downloads_dir, out_filename)
        if os.path.exists(out_path):
            try: os.remove(out_path)
            except: pass

        tpl_path = os.path.join(base_dir, "附件", "现场会照片.docx")
        if not os.path.exists(tpl_path):
            tpl_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "附件", "现场会照片.docx")

        if os.path.exists(tpl_path):
            shutil.copy(tpl_path, out_path)
            doc = word.Documents.Open(out_path)
        else:
            # 模板不存在时的后备容错：新建文档
            doc = word.Documents.Add()
            doc.PageSetup.TopMargin = 72
            doc.PageSetup.BottomMargin = 72
            doc.PageSetup.LeftMargin = 72
            doc.PageSetup.RightMargin = 72

        now_dt = datetime.datetime.now()
        now_date_str = f"{now_dt.year}年{now_dt.month:02d}月{now_dt.day:02d}日"
        target_loc = f"{clean_ts}{clean_vn}" if (clean_ts or clean_vn) else "现场会"

        # 1. 替换标题与元数据中的行政区划与日期
        rng = doc.Content
        # 替换镇村名称
        rng.Find.Execute("襄河镇八波村", False, False, False, False, False, True, 1, False, target_loc, 2)
        # 替换县名（若非全椒县）
        if county_name and county_name != "全椒县":
            rng.Find.Execute("全椒县", False, False, False, False, False, True, 1, False, county_name, 2)
        # 替换拍摄时间
        rng.Find.Execute("拍摄时间：[0-9]{4}年[0-9]{2}月[0-9]{2}日", False, False, True, False, False, True, 1, False, f"拍摄时间：{now_date_str}", 2)

        # 2. 筛选有效照片路径
        valid_photos = []
        for p_item in (photos or []):
            raw_url = p_item if isinstance(p_item, str) else (p_item.get("url") or "")
            if not raw_url:
                continue
            clean_url = raw_url.strip()
            if clean_url.startswith("/uploads/"):
                rel_path = clean_url.lstrip("/")
            elif "file=" in clean_url:
                rel_path = clean_url.split("file=")[-1].lstrip("/")
            else:
                rel_path = os.path.join("uploads", os.path.basename(clean_url))

            candidates = [
                os.path.join(base_dir, "backend", rel_path),
                os.path.join(base_dir, rel_path),
                os.path.join(os.path.dirname(__file__), rel_path),
                os.path.abspath(clean_url)
            ]
            for cand in candidates:
                if os.path.exists(cand) and os.path.isfile(cand):
                    valid_photos.append(cand)
                    break

        # 3. 将照片填充至表格单元格中
        if doc.Tables.Count > 0:
            table = doc.Tables(1)
            table.AllowAutoFit = False  # 禁止表格自动拉伸以保持网格整齐

            if not valid_photos:
                cell = table.Cell(1, 1)
                p = cell.Range.Paragraphs(1)
                p.Range.Text = "（注：当前行政村尚未上传现场会核查照片）"
                p.Range.Font.Name = "宋体"
                p.Range.Font.Size = 11
                p.Range.Font.Color = 8421504  # 灰色
                p.Range.ParagraphFormat.Alignment = 1
                cell.VerticalAlignment = 1
            else:
                # 动态扩展行数：每行容纳 2 张照片
                needed_rows = (len(valid_photos) + 1) // 2
                while table.Rows.Count < needed_rows:
                    new_row = table.Rows.Add()
                    new_row.Height = table.Rows(1).Height
                    new_row.HeightRule = 1  # wdRowHeightAtLeast

                for p_idx, photo_abs in enumerate(valid_photos):
                    r = (p_idx // 2) + 1
                    c = (p_idx % 2) + 1
                    cell = table.Cell(r, c)

                    # 清除可能多余的段落，保留单个居中段落
                    p = cell.Range.Paragraphs(1)
                    p.Range.Text = ""
                    p.Range.ParagraphFormat.SpaceBefore = 0
                    p.Range.ParagraphFormat.SpaceAfter = 0
                    p.Range.ParagraphFormat.Alignment = 1  # 水平居中
                    cell.VerticalAlignment = 1  # 垂直居中

                    try:
                        with Image.open(photo_abs) as im:
                            im = ImageOps.exif_transpose(im)
                            img_w, img_h = im.size

                        # 单元格可用最大宽高（扣除边距与安全余量，避免溢出引发分页）
                        avail_w = max(cell.Width - cell.LeftPadding - cell.RightPadding - 4.0, 10.0)
                        avail_h = max(table.Rows(r).Height - cell.TopPadding - cell.BottomPadding - 8.0, 10.0)

                        scale = min(avail_w / max(img_w, 1), avail_h / max(img_h, 1))
                        target_w = img_w * scale
                        target_h = img_h * scale

                        shape = cell.Range.InlineShapes.AddPicture(
                            FileName=photo_abs,
                            LinkToFile=False,
                            SaveWithDocument=True
                        )
                        shape.LockAspectRatio = -1  # msoTrue
                        shape.Width = target_w
                        shape.Height = target_h
                    except Exception as pe:
                        print(f"插入现场照片 {p_idx + 1} 异常: {pe}")
                        cell.Range.Text = "【照片加载失败】"

            # 压缩表格后末尾段落高度，确保 1~4 张照片时严格保持为单页 A4
            if doc.Paragraphs.Count > 0:
                last_p = doc.Paragraphs(doc.Paragraphs.Count)
                last_p.Range.Font.Size = 1
                last_p.Range.ParagraphFormat.SpaceBefore = 0
                last_p.Range.ParagraphFormat.SpaceAfter = 0
                last_p.Range.ParagraphFormat.LineSpacingRule = 4  # wdLineSpaceExactly
                last_p.Range.ParagraphFormat.LineSpacing = 1
        else:
            # 容错：如果文档无表格则简单追加
            for p_idx, photo_abs in enumerate(valid_photos, 1):
                p_img = doc.Paragraphs.Add()
                shape = p_img.Range.InlineShapes.AddPicture(
                    FileName=photo_abs,
                    LinkToFile=False,
                    SaveWithDocument=True
                )
                shape.Width = 360
                shape.Height = shape.Height * (360.0 / max(shape.Width, 1))

        doc.SaveAs2(FileName=out_path, FileFormat=12)
        doc.Close(0)
        doc = None
        word.Quit()
        word = None
        return f"/api/download?file=downloads/{out_filename}"
    except Exception as e:
        print("export_village_meeting_photos error:", e)
        if doc:
            try: doc.Close(0)
            except: pass
        if word:
            try: word.Quit()
            except: pass
        raise e
    finally:
        pythoncom.CoUninitialize()


def export_sample_detail_excel(township_name: str, contractor_rows: list) -> str:
    """
    导出各乡镇自查抽样农户明细表 Excel 文件 (.xlsx)
    包含字段：序号、乡镇名、村名、组名、承包方编码、承包方名称、承包方证件号码。
    严格遵循 Excel 规范：居中对齐、文本自动换行、单元格全边框、数字与身份证文本格式防变形。
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    base_dir = get_base_dir()
    clean_ts = sanitize_filename(township_name)
    downloads_dir = os.path.join(base_dir, "backend", "downloads")
    os.makedirs(downloads_dir, exist_ok=True)
    out_path = os.path.join(downloads_dir, f"自查抽样明细表_{clean_ts}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "抽样农户明细"
    ws.views.sheetView[0].showGridLines = True

    # 样式定义
    font_title = Font(name="微软雅黑", size=15, bold=True, color="1F497D")
    font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="微软雅黑", size=10)
    font_total = Font(name="微软雅黑", size=10, bold=True, color="000000")

    fill_header = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="BFBFBF")
    double_border_side = Side(border_style="double", color="000000")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_border_side)

    # 1. 标题行
    ws.merge_cells("A1:G1")
    ws["A1"] = f"【{township_name}】二轮延包自查抽样农户明细表"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 38

    # 2. 表头行
    headers = ["序号", "乡镇名", "村名", "组名", "承包方编码", "承包方名称", "承包方证件号码"]
    ws.append(headers)
    ws.row_dimensions[2].height = 26

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=2, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = align_center
        c.border = border_all

    # 3. 填充数据行
    start_row = 3
    villages_set = set()
    groups_set = set()

    for i, r in enumerate(contractor_rows, 1):
        curr_row = start_row + i - 1
        ws.row_dimensions[curr_row].height = 22

        t_name = str(r.get("township_name") or township_name or "").strip()
        v_name = str(r.get("village_name") or "").strip()
        g_name = str(r.get("group_name") or "").strip()
        cbfbm = str(r.get("cbfbm") or "").strip()
        cbfmc = str(r.get("cbfmc") or "").strip()
        cbfzjhm = str(r.get("cbfzjhm") or "").strip()

        if v_name:
            villages_set.add(v_name)
        if g_name:
            groups_set.add((v_name, g_name))

        ws.append([i, t_name, v_name, g_name, cbfbm, cbfmc, cbfzjhm])

        is_zebra = (i % 2 == 0)
        row_fill = fill_zebra if is_zebra else None

        for c_idx in range(1, 8):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.font = font_data
            cell.alignment = align_center
            cell.border = border_all
            if row_fill:
                cell.fill = row_fill

        # 确保编码与身份证号码为严格文本格式防变形
        ws.cell(row=curr_row, column=5).number_format = "@"
        ws.cell(row=curr_row, column=7).number_format = "@"

    # 4. 统计汇总行
    total_row = start_row + len(contractor_rows)
    ws.row_dimensions[total_row].height = 26
    total_data = [
        "合计",
        township_name,
        f"共 {len(villages_set)} 个行政村",
        f"共 {len(groups_set)} 个组",
        f"总抽样: {len(contractor_rows)} 户",
        "-",
        "-"
    ]
    ws.append(total_data)

    for c_idx in range(1, 8):
        cell = ws.cell(row=total_row, column=c_idx)
        cell.font = font_total
        cell.fill = fill_total
        cell.alignment = align_center
        cell.border = border_total

    # 5. 列宽自适应
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row == 1:
                continue
            val_str = str(cell.value or "")
            str_len = sum(2 if ord(char) > 127 else 1 for char in val_str)
            if str_len > max_len:
                max_len = str_len
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # 专门优化各列宽度
    ws.column_dimensions["A"].width = 8   # 序号
    ws.column_dimensions["B"].width = 14  # 乡镇名
    ws.column_dimensions["C"].width = 16  # 村名
    ws.column_dimensions["D"].width = 16  # 组名
    ws.column_dimensions["E"].width = 24  # 承包方编码 (18位)
    ws.column_dimensions["F"].width = 15  # 承包方名称
    ws.column_dimensions["G"].width = 24  # 承包方证件号码 (18位)

    wb.save(out_path)
    return f"/api/download?file=downloads/自查抽样明细表_{clean_ts}.xlsx"

