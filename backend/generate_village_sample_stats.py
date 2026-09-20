import asyncio
import os
import math
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from database import SessionLocal, get_county_and_townships_sync

async def query_village_sample_stats():
    """
    从数据库精准查询各行政村抽样统计数据
    """
    async with SessionLocal() as session:
        # 查询所有涉及的外业抽样村
        sql = """
            SELECT 
                township_name,
                village_name,
                SUBSTRING(group_code, 1, 12) || '00' as village_code,
                SUBSTRING(group_code, 1, 12) as village_prefix,
                COUNT(DISTINCT group_code) as group_cnt,
                COUNT(*) as parcel_cnt,
                COUNT(DISTINCT cbfbm) as sample_cbf_cnt
            FROM waiye_samples
            WHERE township_name IS NOT NULL AND village_name IS NOT NULL
            GROUP BY township_name, village_name, SUBSTRING(group_code, 1, 12)
            ORDER BY township_name, village_code
        """
        res = await session.execute(text(sql))
        rows = res.fetchall()

        village_list = []
        for r in rows:
            t_name, v_name, v_code, v_prefix, g_cnt, p_cnt, s_cbf = r
            # 查询该村在 cbf 表中的全村总农户数
            res_cbf = await session.execute(
                text("SELECT COUNT(*) FROM cbf WHERE cbfbm::text LIKE :code"),
                {"code": f"{v_prefix}%"}
            )
            total_cbf = res_cbf.scalar() or 0
            rate = (s_cbf / total_cbf) if total_cbf > 0 else 0
            req_5 = math.ceil(total_cbf * 0.05)
            is_ok = (s_cbf >= req_5)
            if is_ok:
                status = "达标 (≥5%)"
            else:
                shortage = req_5 - s_cbf
                status = f"未达标 (差{shortage}户)"

            village_list.append({
                "township": t_name,
                "village": v_name,
                "village_code": v_code,
                "group_cnt": g_cnt,
                "parcel_cnt": p_cnt,
                "sample_cbf": s_cbf,
                "total_cbf": total_cbf,
                "rate": rate,
                "status": status,
                "is_ok": is_ok,
                "required_5pct": req_5
            })

        return village_list

def build_excel(village_data, output_paths):
    """
    创建带边框、单元格居中、开启自动换行的标准 Excel 统计表
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "各村抽样比例统计"
    ws.views.sheetView[0].showGridLines = True

    # 样式定义
    font_title = Font(name="微软雅黑", size=16, bold=True, color="1F497D")
    font_sub = Font(name="微软雅黑", size=9.5, italic=False, color="595959")
    font_header = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="微软雅黑", size=10, bold=False, color="000000")
    font_total = Font(name="微软雅黑", size=11, bold=True, color="000000")
    
    font_status_ok = Font(name="微软雅黑", size=10, bold=True, color="274E13")
    fill_status_ok = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    font_status_warn = Font(name="微软雅黑", size=10, bold=True, color="9C0006")
    fill_status_warn = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    fill_header = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    fill_zebra = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_data = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    
    header_border_side = Side(border_style="thin", color="1B4F72")
    border_header = Border(left=header_border_side, right=header_border_side, top=header_border_side, bottom=header_border_side)

    total_border_top = Side(border_style="thin", color="2E75B6")
    total_border_bottom = Side(border_style="double", color="2E75B6")
    border_total = Border(left=thin_border_side, right=thin_border_side, top=total_border_top, bottom=total_border_bottom)

    # 居中对齐并开启自动换行 (AGENTS.md 要求)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. 标题行 (Row 1)
    ws.merge_cells("A1:J1")
    cell_a1 = ws["A1"]
    cell_a1.value = "全椒县二轮延包县级验收 — 各行政村抽样比例统计表"
    cell_a1.font = font_title
    cell_a1.alignment = align_center
    ws.row_dimensions[1].height = 40

    # 2. 说明行 (Row 2)
    ws.merge_cells("A2:J2")
    cell_a2 = ws["A2"]
    now_str = datetime.datetime.now().strftime("%Y年%m月%d日")
    cell_a2.value = f"统计口径：按行政村统计外业实际抽查农户数占全村总农户数比例 | 达标标准：实际抽样农户数 ≥ 全村农户总数 × 5% (向上取整) | 统计时间：{now_str}"
    cell_a2.font = font_sub
    cell_a2.alignment = align_center
    ws.row_dimensions[2].height = 24

    # 3. 表头行 (Row 3)
    headers = [
        "序号", "所属乡镇", "行政村名称", "行政村代码", 
        "抽样组数", "抽样地块数", "实际抽样农户数", "全村总农户数", 
        "抽样占比", "5%达标状态"
    ]
    ws.row_dimensions[3].height = 30
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header

    # 4. 数据行 (Row 4 开始)
    cur_row = 4
    for idx, item in enumerate(village_data, 1):
        ws.row_dimensions[cur_row].height = 24
        
        row_values = [
            idx,
            item["township"],
            item["village"],
            str(item["village_code"]),
            item["group_cnt"],
            item["parcel_cnt"],
            item["sample_cbf"],
            item["total_cbf"],
            item["rate"],
            item["status"]
        ]

        # 隔行轻微斑马纹底色
        row_fill = fill_zebra if idx % 2 == 0 else None

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=cur_row, column=col_idx, value=val)
            cell.font = font_data
            cell.alignment = align_center
            cell.border = border_data
            if row_fill:
                cell.fill = row_fill

            # 抽样占比列设置为百分比格式 0.00%
            if col_idx == 9:
                cell.number_format = "0.00%"

            # 达标状态列高亮
            if col_idx == 10:
                if item["is_ok"]:
                    cell.font = font_status_ok
                    cell.fill = fill_status_ok
                else:
                    cell.font = font_status_warn
                    cell.fill = fill_status_warn

        cur_row += 1

    # 5. 合计汇总行
    ws.row_dimensions[cur_row].height = 28
    tot_townships = len(set(d["township"] for d in village_data))
    tot_villages = len(village_data)
    tot_groups = sum(d["group_cnt"] for d in village_data)
    tot_parcels = sum(d["parcel_cnt"] for d in village_data)
    tot_samples = sum(d["sample_cbf"] for d in village_data)
    tot_cbf = sum(d["total_cbf"] for d in village_data)
    tot_rate = (tot_samples / tot_cbf) if tot_cbf > 0 else 0
    
    ok_count = sum(1 for d in village_data if d["is_ok"])
    if ok_count == tot_villages:
        tot_status = f"{tot_villages}村全部达标 (100%)"
    else:
        tot_status = f"{ok_count}村达标 / {tot_villages - ok_count}村未达标"

    total_values = [
        "合计",
        f"覆盖{tot_townships}个乡镇",
        f"共{tot_villages}个村",
        "-",
        tot_groups,
        tot_parcels,
        tot_samples,
        tot_cbf,
        tot_rate,
        tot_status
    ]

    for col_idx, val in enumerate(total_values, 1):
        cell = ws.cell(row=cur_row, column=col_idx, value=val)
        cell.font = font_total
        cell.fill = fill_total
        cell.alignment = align_center
        cell.border = border_total
        if col_idx == 9:
            cell.number_format = "0.00%"
        if col_idx == 10:
            if ok_count == tot_villages:
                cell.font = font_status_ok
            else:
                cell.font = font_status_warn

    # 6. 列宽自适应与优化
    col_widths = {
        "A": 9.0,   # 序号
        "B": 14.0,  # 所属乡镇
        "C": 16.0,  # 行政村名称
        "D": 19.0,  # 行政村代码
        "E": 12.0,  # 抽样组数
        "F": 14.0,  # 抽样地块数
        "G": 18.0,  # 实际抽样农户数
        "H": 16.0,  # 全村总农户数
        "I": 14.0,  # 抽样占比
        "J": 24.0   # 5%达标状态
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # 冻结前三行（滚动时表头固定可见）
    ws.freeze_panes = "A4"

    # 保存到所有指定路径
    for p in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(p)), exist_ok=True)
        wb.save(p)
        print(f"Successfully saved: {p}")

async def main():
    print("=== 开始从数据库重新统计全椒县各行政村抽样比例统计表 ===")
    village_data = await query_village_sample_stats()
    print(f"统计完成，共获取 {len(village_data)} 个抽样行政村的数据。")
    
    root_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "全椒县各行政村抽样比例统计表.xlsx"))
    download_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "downloads", "全椒县各行政村抽样比例统计表.xlsx"))
    
    build_excel(village_data, [root_path, download_path])
    print("=== 全部生成完成！ ===")

if __name__ == "__main__":
    asyncio.run(main())
